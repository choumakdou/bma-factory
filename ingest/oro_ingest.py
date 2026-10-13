"""ORO case ingestion — registry-driven extraction into a queryable database.

Why this works at all: the master workbook's `MailMerge` sheet is not a list of
merge placeholders, it is a *declared mapping table*:

    # | Field Name | Value (=source formula) | Source Sheet | Description

Every one of the 88 rows says exactly which sheet and cell holds the value. So
ingestion is not heuristic scraping — it is resolving a mapping that already
exists. Any case workbook that follows the master's sheet structure parses
without new code, which is what makes ~400 reports tractable.

The design principle is the same as the generator: **do not trust the parse,
prove it.** A field that resolves is not necessarily a field that resolved
*correctly*. `check()` below is a type and plausibility gate, and a case that
fails it is recorded with its problems rather than silently imported.

Usage:
    python3 -m ingest.oro_ingest --master <master.xlsx> <case.xlsx> [...]
    python3 -m ingest.oro_ingest --registry <master.xlsx>     # dump the schema
"""

from __future__ import annotations

import argparse
import re
import sqlite3
import sys
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

REF_RE = re.compile(r"=?'?([^'!]+)'?!\$?([A-Z]{1,3})\$?(\d+)")

# --- field typing -----------------------------------------------------------
# Not every registry field is free text. This table says what a field must look
# like if it is going to be trusted, and is the actual QC gate on ingestion.
TYPES: dict[str, str] = {
    "Fee_HKD": "number",
    "SA_SqFt": "number",
    "GFA_SqFt": "number",
    "Lot_Shares": "text",              # "60/414,930" — not a ratio we can cast
    "Floor": "int",
    "Flat": "text",                    # may be "17A"
    "HKHA_Flat_No": "text",
    "Year_of_Completion": "int",
    "Land_Lease_Term_Years": "int",
    "HOS_Initial_Market_Value": "number",
    "HOS_Market_Value_After_LegCO_Extension": "number",
    "ORO_Letter_Date": "date",
    "ORO_Letter_Deadline": "date",
    "DMC_Date": "date",
    "Valuation_Date": "date",
    "Lot_Number": "lot",
    "District_Code": "district_code",
}

# Fields that are legitimately blank in a template (computed or manual-fill
# standards citations), so blank is not a defect for them.
OPTIONAL_BLANK = {
    "HKIS_Standards_Cited", "RICS_Standards_Cited", "IVS_Cited",
    "HOS_Market_Value_After_LegCO_Extension",
    "Owner_2_English", "Owner_2_Chinese", "Lot_Shares",
}

DISTRICT_CODES = {
    "HK", "KC", "KW", "KT", "TWD", "TWN", "TPE",
    "HKC", "HKT", "SK", "ST", "TP", "TWN", "TWD", "TWP",
    "NT", "TKO", "TW", "Yuen", "YK", "YL",
}

LOT_RE = re.compile(r"(Town\s+Lot|Lot\s+No\.?|Kowloon\s+Lot|Inland\s+Lot)", re.IGNORECASE)


@dataclass
class FieldDef:
    index: int
    name: str
    formula: str
    source_sheet: str = ""
    source_cell: str = ""


@dataclass
class Problem:
    severity: str   # "error" | "warning"
    field: str
    message: str

    def __str__(self) -> str:
        icon = {"error": "🔴", "warning": "🟡"}[self.severity]
        return f"{icon} {self.field}: {self.message}"


@dataclass
class ParsedCase:
    values: dict[str, Any] = field(default_factory=dict)
    problems: list[Problem] = field(default_factory=list)
    source_file: str = ""
    resolved: int = 0
    unresolved: list[str] = field(default_factory=list)


# --- registry ---------------------------------------------------------------

def read_registry(master_path: Path) -> list[FieldDef]:
    wb = load_workbook(master_path, data_only=False)
    ws = wb["MailMerge"]
    defs: list[FieldDef] = []
    for r in range(5, ws.max_row + 1):
        idx = ws.cell(row=r, column=2).value
        name = ws.cell(row=r, column=3).value
        formula = ws.cell(row=r, column=4).value
        # rows headed "Notes" etc. are not fields
        if name is None or not isinstance(idx, (int, float)):
            continue
        s = str(formula) if formula else ""
        m = REF_RE.match(s)
        defs.append(
            FieldDef(
                index=int(idx),
                name=str(name).strip(),
                formula=s,
                source_sheet=m.group(1).strip() if m else "",
                source_cell=f"{m.group(2)}{m.group(3)}" if m else "",
            )
        )
    return defs


def parse_case(case_path: Path, defs: list[FieldDef]) -> ParsedCase:
    """Resolve every registry field against one case workbook."""
    wb = load_workbook(case_path, data_only=False)
    out = ParsedCase(source_file=case_path.name)

    for d in defs:
        if not d.source_sheet:
            out.unresolved.append(d.name)
            continue
        if d.source_sheet not in wb.sheetnames:
            out.problems.append(
                Problem("error", d.name, f"source sheet '{d.source_sheet}' not in case workbook")
            )
            out.unresolved.append(d.name)
            continue
        v = wb[d.source_sheet][d.source_cell].value
        out.values[d.name] = "" if v is None else v
        out.resolved += 1

    _check(out)
    return out


# --- validation -------------------------------------------------------------

def _as_num(v: Any) -> float | None:
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "").replace("HK$", "").replace("$", "")
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _check(case: ParsedCase) -> None:
    """Type and plausibility gate. A resolved field is not a correct field."""
    v = case.values

    for name, kind in TYPES.items():
        raw = v.get(name)
        if raw is None or (isinstance(raw, str) and not raw.strip()):
            if name not in OPTIONAL_BLANK:
                case.problems.append(Problem("warning", name, "empty"))
            continue

        if kind == "number":
            n = _as_num(raw)
            if n is None:
                # resolved to prose where a number is required — the single most
                # important thing this gate catches
                case.problems.append(
                    Problem("error", name, f"expected a number, got {str(raw)[:60]!r}")
                )
            elif n <= 0:
                case.problems.append(Problem("error", name, f"non-positive value {n:g}"))
        elif kind == "int":
            n = _as_num(raw)
            if n is None or n != int(n):
                case.problems.append(Problem("warning", name, f"expected a whole number, got {raw!r}"))
        elif kind == "date":
            if not isinstance(raw, (date, datetime)) and not re.match(r"^\d{4}-\d{2}-\d{2}", str(raw)):
                case.problems.append(Problem("warning", name, f"expected a date, got {str(raw)[:40]!r}"))
        elif kind == "district_code":
            if str(raw).strip() not in DISTRICT_CODES:
                case.problems.append(Problem("warning", name, f"unrecognised district code {raw!r}"))
        elif kind == "lot":
            if not LOT_RE.search(str(raw)):
                case.problems.append(Problem("warning", name, f"does not look like a lot: {str(raw)[:40]!r}"))

    # cross-field checks
    if not v.get("ORO_File_Ref"):
        case.problems.append(Problem("error", "ORO_File_Ref", "missing — cannot identify the case"))
    if not v.get("Property_Address_English"):
        case.problems.append(Problem("error", "Property_Address_English", "missing — cannot identify the property"))
    if not v.get("Lot_Number"):
        case.problems.append(Problem("error", "Lot_Number", "missing — property cannot be deduplicated"))

    # the Chinese address must actually be in Chinese; the English one must
    # not be, or the two have been swapped somewhere upstream
    zh = str(v.get("Property_Address_Chinese") or "")
    en = str(v.get("Property_Address_English") or "")
    if zh and not re.search(r"[\u4e00-\u9fff]", zh):
        case.problems.append(Problem("warning", "Property_Address_Chinese", "no Chinese characters found"))
    if en and re.search(r"[\u4e00-\u9fff]", en):
        case.problems.append(Problem("warning", "Property_Address_English", "contains Chinese — check EN/ZH are not swapped"))


# --- database ---------------------------------------------------------------

SCHEMA_SQL = """
PRAGMA foreign_keys = ON;

CREATE TABLE IF NOT EXISTS properties (
    id            INTEGER PRIMARY KEY,
    address_en    TEXT, address_zh   TEXT,
    estate_en     TEXT, estate_zh    TEXT,
    lot_number    TEXT NOT NULL,
    block_en      TEXT, block_zh     TEXT,
    floor         TEXT, flat         TEXT,
    district      TEXT, district_code TEXT,
    property_type TEXT,
    year_completion INTEGER,
    sa_sqft       REAL, gfa_sqft     REAL,
    land_lease_years INTEGER, land_lease_commencement TEXT,
    hos_initial_market_value REAL, hos_market_value_legco REAL,
    lat REAL, lng REAL,
    first_seen    TEXT, last_seen    TEXT,
    case_count    INTEGER DEFAULT 0,
    UNIQUE (lot_number, floor, flat)
);

CREATE TABLE IF NOT EXISTS cases (
    id            INTEGER PRIMARY KEY,
    property_id   INTEGER NOT NULL REFERENCES properties(id),
    oro_file_ref  TEXT UNIQUE,
    chtf_ref      TEXT,
    letter_date   TEXT, deadline     TEXT, valuation_date TEXT,
    bankruptcy_no TEXT,
    bankrupt_en   TEXT, bankrupt_zh  TEXT,
    valuer        TEXT, fee_hkd      REAL,
    source_file   TEXT, source_hash  TEXT,
    ingested_at   TEXT
);

CREATE TABLE IF NOT EXISTS case_owners (
    case_id INTEGER REFERENCES cases(id) ON DELETE CASCADE,
    ord     INTEGER, name_en TEXT, name_zh TEXT, capacity TEXT
);

CREATE TABLE IF NOT EXISTS case_encumbrances (
    case_id INTEGER REFERENCES cases(id) ON DELETE CASCADE,
    ord     INTEGER, memorial TEXT, enc_date TEXT,
    nature  TEXT, parties TEXT, consideration TEXT
);

CREATE TABLE IF NOT EXISTS case_raw (
    case_id INTEGER REFERENCES cases(id) ON DELETE CASCADE,
    field_name TEXT, value TEXT
);

CREATE TABLE IF NOT EXISTS ingest_problems (
    case_id INTEGER REFERENCES cases(id) ON DELETE CASCADE,
    severity TEXT, field TEXT, message TEXT
);

CREATE INDEX IF NOT EXISTS ix_cases_property ON cases(property_id);
CREATE INDEX IF NOT EXISTS ix_cases_date     ON cases(valuation_date);
CREATE INDEX IF NOT EXISTS ix_prop_district  ON properties(district_code);
"""


def _txt(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, (date, datetime)):
        return v.isoformat()[:10]
    return str(v).strip()


def _num(v: Any) -> float | None:
    return _as_num(v)


def connect(db_path: Path) -> sqlite3.Connection:
    db_path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(db_path)
    conn.executescript(SCHEMA_SQL)
    return conn


def upsert_case(conn: sqlite3.Connection, case: ParsedCase, strict: bool = False) -> int | None:
    """Insert one parsed case. Returns the case id, or None if refused."""
    errs = [p for p in case.problems if p.severity == "error"]
    if errs and strict:
        return None

    v = case.values
    now = datetime.now().isoformat(timespec="seconds")
    val_date = _txt(v.get("Valuation_Date")) or _txt(v.get("ORO_Letter_Date")) or now[:10]

    # --- property (dedupe on lot + floor + flat) ---
    key = (_txt(v.get("Lot_Number")), _txt(v.get("Floor")), _txt(v.get("Flat")))
    cur = conn.execute(
        "SELECT id FROM properties WHERE lot_number=? AND IFNULL(floor,'')=? AND IFNULL(flat,'')=?",
        key,
    )
    row = cur.fetchone()
    if row:
        pid = row[0]
        conn.execute(
            "UPDATE properties SET last_seen=?, case_count=case_count+1,"
            " sa_sqft=COALESCE(?,sa_sqft), gfa_sqft=COALESCE(?,gfa_sqft) WHERE id=?",
            (val_date, _num(v.get("SA_SqFt")), _num(v.get("GFA_SqFt")), pid),
        )
    else:
        cur = conn.execute(
            """INSERT INTO properties
               (address_en,address_zh,estate_en,estate_zh,lot_number,block_en,block_zh,
                floor,flat,district,district_code,property_type,year_completion,
                sa_sqft,gfa_sqft,land_lease_years,land_lease_commencement,
                hos_initial_market_value,hos_market_value_legco,
                first_seen,last_seen,case_count)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,1)""",
            (_txt(v.get("Property_Address_English")), _txt(v.get("Property_Address_Chinese")),
             _txt(v.get("Estate_Name_English")), _txt(v.get("Estate_Name_Chinese")),
             key[0], _txt(v.get("Block_English")), _txt(v.get("Block_Chinese")),
             key[1], key[2], _txt(v.get("District")), _txt(v.get("District_Code")),
             _txt(v.get("Property_Type")), _num(v.get("Year_of_Completion")),
             _num(v.get("SA_SqFt")), _num(v.get("GFA_SqFt")),
             _num(v.get("Land_Lease_Term_Years")), _txt(v.get("Land_Lease_Commencement")),
             _num(v.get("HOS_Initial_Market_Value")),
             _num(v.get("HOS_Market_Value_After_LegCO_Extension")),
             val_date, val_date),
        )
        pid = cur.lastrowid

    cur = conn.execute(
        """INSERT OR REPLACE INTO cases
           (property_id,oro_file_ref,chtf_ref,letter_date,deadline,valuation_date,
            bankruptcy_no,bankrupt_en,bankrupt_zh,valuer,fee_hkd,source_file,ingested_at)
           VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        (pid, _txt(v.get("ORO_File_Ref")), _txt(v.get("CHFT_Ref")),
         _txt(v.get("ORO_Letter_Date")), _txt(v.get("ORO_Letter_Deadline")), val_date,
         _txt(v.get("Bankruptcy_No")), _txt(v.get("Bankrupt_English")),
         _txt(v.get("Bankrupt_Chinese")), _txt(v.get("Valuer_Name")),
         _num(v.get("Fee_HKD")), case.source_file, now),
    )
    cid = cur.lastrowid

    conn.execute("DELETE FROM case_owners WHERE case_id=?", (cid,))
    conn.execute("DELETE FROM case_encumbrances WHERE case_id=?", (cid,))
    conn.execute("DELETE FROM case_raw WHERE case_id=?", (cid,))
    conn.execute("DELETE FROM ingest_problems WHERE case_id=?", (cid,))

    for i in (1, 2):
        nm = _txt(v.get(f"Owner_{i}_English")) or _txt(v.get(f"Owner_{i}_Chinese"))
        if nm:
            conn.execute(
                "INSERT INTO case_owners (case_id,ord,name_en,name_zh,capacity) VALUES (?,?,?,?,?)",
                (cid, i, _txt(v.get(f"Owner_{i}_English")), _txt(v.get(f"Owner_{i}_Chinese")),
                 _txt(v.get("Capacity"))),
            )

    for i in range(1, 9):
        mem = _txt(v.get(f"Encumbrance_{i}_Memorial"))
        if not mem:
            continue
        conn.execute(
            "INSERT INTO case_encumbrances (case_id,ord,memorial,enc_date,nature,parties,consideration)"
            " VALUES (?,?,?,?,?,?,?)",
            (cid, i, mem, _txt(v.get(f"Encumbrance_{i}_Date")), _txt(v.get(f"Encumbrance_{i}_Nature")),
             _txt(v.get(f"Encumbrance_{i}_Parties")), _txt(v.get(f"Encumbrance_{i}_Consideration"))),
        )

    conn.executemany(
        "INSERT INTO case_raw (case_id,field_name,value) VALUES (?,?,?)",
        [(cid, k, _txt(x)) for k, x in v.items()],
    )
    conn.executemany(
        "INSERT INTO ingest_problems (case_id,severity,field,message) VALUES (?,?,?,?)",
        [(cid, p.severity, p.field, p.message) for p in case.problems],
    )
    conn.commit()
    return cid


# --- reporting --------------------------------------------------------------

def report(case: ParsedCase) -> None:
    errs = [p for p in case.problems if p.severity == "error"]
    warns = [p for p in case.problems if p.severity == "warning"]
    print(f"\n{case.source_file}")
    print(f"  resolved {case.resolved} fields, {len(case.unresolved)} unresolved")
    for p in errs:
        print(f"  {p}")
    for p in warns[:8]:
        print(f"  {p}")
    if len(warns) > 8:
        print(f"  … {len(warns) - 8} more warnings")
    if not errs and not warns:
        print("  ✅ clean — all types and cross-checks pass")


def main() -> int:
    ap = argparse.ArgumentParser(description="Ingest ORO case workbooks into SQLite")
    ap.add_argument("cases", nargs="*", type=Path, help="case workbooks (.xlsx)")
    ap.add_argument("--master", type=Path, required=True, help="master workbook containing the MailMerge registry")
    ap.add_argument("--db", type=Path, default=Path("data/ORO Cases.sqlite3"))
    ap.add_argument("--registry", action="store_true", help="dump the field registry and exit")
    ap.add_argument("--strict", action="store_true", help="refuse cases with validation errors")
    args = ap.parse_args()

    defs = read_registry(args.master)

    if args.registry:
        print(f"{len(defs)} fields\n")
        for d in defs:
            t = TYPES.get(d.name, "")
            print(f"  {d.index:3d}  {d.name:42s} {d.source_sheet:20s}!{d.source_cell:5s} {t}")
        return 0

    if not args.cases:
        ap.error("supply at least one case workbook, or use --registry")

    conn = connect(args.db)
    n_ok = 0
    for path in args.cases:
        try:
            case = parse_case(path, defs)
        except Exception as exc:                      # a broken file must not stop the run
            print(f"\n{path.name}\n  🔴 failed to open: {exc}")
            continue
        report(case)
        cid = upsert_case(conn, case, strict=args.strict)
        if cid is None:
            print("  ⛔ refused (strict mode)")
        else:
            n_ok += 1
            print(f"  → case_id {cid}")
    print(f"\ningested {n_ok}/{len(args.cases)} → {args.db}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
